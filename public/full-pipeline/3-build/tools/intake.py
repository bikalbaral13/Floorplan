"""S1 intake: plan-phase deck (.pptx or .pdf) -> deck.json + extracted images (+ pages/ renders for a PDF).

Reads every slide with python-pptx and records text (with position), links,
dimensions, pictures (deduplicated by md5, with crop) and drawn overlay shapes.
Then segments the deck into projects -> rooms and classifies each slide.
No guessing is written as fact: every item keeps its slide number.

    python tools/intake.py <deck.pptx> <out_dir>
"""
import hashlib, json, re, sys
from collections import Counter
from pathlib import Path
from pptx import Presentation
from pptx.util import Emu

CATEGORIES = [  # checked against the first line of each text box
    ("ceiling", r"ceiling|baffle"),
    ("flooring", r"floor|carpet|lvt|rug"),
    ("wall", r"wall|partition|panel|column|glass|door|graphics|logo|paint|texture|moss|cladding"),
    ("lighting", r"light|pendant|lamp"),
    ("furniture", r"furniture|sofa|chair|table|workstation|storage|seating|stool|pantry|printer|planter|pouf"),
]
DIM = re.compile(r"(DIA\.?\s*)?(\d{2,5})\s*(mm|m)\b", re.I)
LEVEL = re.compile(r"(\d{3,4})\s*mm\s*(lvl|level)", re.I)
URL = re.compile(r"(https?://\S+|\b[\w-]+\.(?:com|co\.uk|it|in|ua)\b\S*)", re.I)
COVER = re.compile(r"(-\s*20\d\d|\(competition\)|\d{2,}\s*seater)", re.I)
# a review comment starts with an instruction, or criticises a render anywhere in the sentence
REVIEW_START = re.compile(r"^\s*(change|add|increase|decrease|reduce|remove|make|move|tone down|use|show|please|"
                          r"another view|the .{0,40}\b(looks|is placed)|this view|all the|services look|lighting to be)\b",
                          re.I | re.M)
REVIEW_ANY = re.compile(r"\b(looks? (too|to)|pixelat|overlap|too low|too high|make sure|to be added|to be adjusted|"
                        r"to be brought|to be as same|need to be|needs to be|be incorporated)\b", re.I)


def is_review(texts):
    return any(REVIEW_START.search(t) or REVIEW_ANY.search(t) for t in texts)


def norm(s):
    s = re.sub(r"[^a-z0-9]", "", s.lower()).replace("colab", "collab")
    return re.sub(r"0(\d)$", r"\1", s)  # collab01 == collab1


def emu_box(sh, W, H):
    return [round(sh.left / W, 4), round(sh.top / H, 4), round(sh.width / W, 4), round(sh.height / H, 4)]


def shape_texts(sh):
    if not sh.has_text_frame:
        return "", []
    links = [r.hyperlink.address for p in sh.text_frame.paragraphs for r in p.runs if r.hyperlink and r.hyperlink.address]
    return sh.text_frame.text.strip(), links


def walk(shapes):
    for sh in shapes:
        if sh.shape_type == 6:  # group
            yield from walk(sh.shapes)
        else:
            yield sh


A_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def freeform_points(sh, W, H):
    """Vertices of a freeform / line in slide-normalised coords (custGeom path space -> shape box).
    Curves keep only their end points, enough for zone outlines."""
    paths = sh._element.findall(f".//{A_NS}custGeom/{A_NS}pathLst/{A_NS}path")
    if not paths and sh.shape_type == 9:
        return [[sh.left / W, sh.top / H], [(sh.left + sh.width) / W, (sh.top + sh.height) / H]]
    out = []
    for path in paths:
        pw, ph = int(path.get("w", 0)) or sh.width, int(path.get("h", 0)) or sh.height
        for cmd in path:
            pts = cmd.findall(f"{A_NS}pt")
            if pts:
                pt = pts[-1]
                x = sh.left + int(pt.get("x")) / pw * sh.width
                y = sh.top + int(pt.get("y")) / ph * sh.height
                out.append([round(x / W, 5), round(y / H, 5)])
    return out


def solid_fill(sh):
    """Swatches in these decks are solid-filled shapes: return '#RRGGBB' (or the theme slot)."""
    try:
        if sh.fill.type == 1:
            c = sh.fill.fore_color
            return f"#{c.rgb}" if c.type == 1 else f"theme:{c.theme_color}"
    except Exception:
        pass
    return None


def read_slide(i, slide, W, H, img_dir, seen):
    rec = {"slide": i, "texts": [], "links": [], "pictures": [], "shapes": [], "tables": []}
    for sh in walk(slide.shapes):
        box = emu_box(sh, W, H)
        if sh.shape_type == 13:
            im = sh.image
            md5 = hashlib.md5(im.blob).hexdigest()
            name = f"{md5[:12]}.{im.ext}"
            if md5 not in seen:
                (img_dir / name).write_bytes(im.blob)
                seen[md5] = {"file": f"images/{name}", "px": list(im.size)}
            crop = [round(sh.crop_left, 5), round(sh.crop_top, 5), round(sh.crop_right, 5), round(sh.crop_bottom, 5)]
            rec["pictures"].append({"md5": md5, "file": seen[md5]["file"], "px": seen[md5]["px"], "box": box, "crop": crop})
            continue
        if sh.has_table:
            rec["tables"].append([[c.text for c in r.cells] for r in sh.table.rows])
            continue
        txt, links = shape_texts(sh)
        rec["links"] += links
        try:
            if sh.click_action.hyperlink.address:
                rec["links"].append(sh.click_action.hyperlink.address)
        except Exception:
            pass
        if txt:
            t = {"text": txt, "box": box}
            fill = solid_fill(sh)
            if fill:
                t["fill"] = fill
            rec["texts"].append(t)
        elif sh.shape_type in (1, 5, 9):  # autoshape / freeform / line drawn over images (callouts)
            kind = None
            try:
                kind = str(sh.auto_shape_type).split(".")[-1].split(" ")[0]
            except Exception:
                kind = {5: "FREEFORM", 9: "LINE"}.get(sh.shape_type)
            shp = {"kind": kind or "SHAPE", "box": box, "name": sh.name}
            fill = solid_fill(sh)
            if fill:
                shp["fill"] = fill
            poly = freeform_points(sh, W, H)
            if poly:
                shp["points"] = poly
            rec["shapes"].append(shp)
    all_text = " \n ".join(t["text"] for t in rec["texts"])
    rec["urls"] = sorted(set(rec["links"]) | {u.rstrip(").,") for u in URL.findall(all_text)})
    rec["dims_mm"] = [("DIA " if d[0] else "") + d[1] + ("000" if d[2].lower() == "m" and not d[0] else "") for d in DIM.findall(all_text)]
    rec["levels_mm"] = [int(m[0]) for m in LEVEL.findall(all_text)]
    del rec["links"]
    return rec


def category(rec):
    for t in rec["texts"]:
        s = t["text"].split("\n")[0]
        if URL.fullmatch(s.strip()):
            continue
        hits = [(m.start(), cat) for cat, rx in CATEGORIES for m in [re.search(rx, s, re.I)] if m]
        if hits:
            return min(hits)[1]  # earliest keyword in the line wins
    return None


def segment(slides):
    """projects -> rooms. Cover slides open a project; a slide whose only text is a room
    name opens that room the first time and marks a final render afterwards."""
    projects, proj, room = [], None, None
    for r in slides:
        texts = [t["text"] for t in r["texts"]]
        cover = next((t for t in texts if COVER.search(t.split("\n")[0])), None)
        if cover and len(r["pictures"]) <= 1 and not r["urls"]:
            rooms = [x.strip() for t in texts if t != cover for x in t.split("\n") if x.strip()]
            proj = {"title": cover.split("\n")[0].strip(), "cover_slide": r["slide"], "room_list": rooms, "rooms": []}
            projects.append(proj)
            room = None
            r["kind"] = "cover"
            if not rooms:  # single-space project (no dividers): one room named after the project
                room = {"name": proj["title"], "id": norm(proj["title"]), "slides": []}
                proj["rooms"].append(room)
            continue
        if proj is None:
            r["kind"] = "front"
            continue
        known = {norm(x): x for x in proj["room_list"]}
        only = texts[0] if len(texts) == 1 else None
        label = norm(only) if only else None
        if label and len(r["pictures"]) <= 1 and not r["urls"] and only.isupper() and len(only) < 30:
            # fuzzy: CONFERENCE == CONFERENCE ROOM, COLAB-01 == COLLAB 01
            match = next((k for k in known if k == label or k.startswith(label) or label.startswith(k)), None)
            existing = next((x for x in proj["rooms"] if x["id"] == (match or label)), None)
            if match and existing is None:
                label = match
                room = {"name": known.get(label, only), "id": label, "divider_slide": r["slide"], "slides": []}
                proj["rooms"].append(room)
                r["kind"] = "divider"
                continue
            r["kind"] = "final_render" if existing is not None else "label"
            r["label"] = only
            if existing is not None and existing is not room:
                r["label_mismatch"] = True
        elif not texts and r["pictures"]:
            r["kind"] = "view"
        elif is_review(texts):
            r["kind"] = "review"
        elif r["dims_mm"] and not r["urls"] and category(r) is None:
            r["kind"] = "dims"
        else:
            r["kind"] = "spec"
        r["category"] = category(r)
        if room is not None:
            room["slides"].append(r["slide"])
        else:
            proj.setdefault("_orphans", []).append(r)
    # Slides that came before any divider (a deck EXTRACT, or a divider placed at the end): give each to the room its
    # first line names ("Collab 1 Ceiling" → COLLAB 01), else — one room only — to that room. Never guessed across rooms.
    for p in projects:
        for r in p.pop("_orphans", []):
            first = norm((r["texts"][0]["text"].split("\n")[0]) if r["texts"] else "")
            hit = next((rm for rm in p["rooms"] if first and first.startswith(rm["id"])), None)
            if hit is None and len(p["rooms"]) == 1:
                hit = p["rooms"][0]
            if hit is None:
                known = {norm(x): x for x in p["room_list"]}
                rid = next((k for k in known if first.startswith(k)), None)
                if rid:
                    hit = {"name": known[rid], "id": rid, "slides": []}
                    p["rooms"].append(hit)
            if hit is not None:
                hit["slides"] = sorted(set(hit["slides"]) | {r["slide"]})
                r["assigned"] = "by title" if len(p["rooms"]) > 1 or first.startswith(hit["id"]) else "only room"
    return projects


def plan_images(projects, by_slide):
    """The floor plan is the picture that recurs on many slides of a project with a crop."""
    for p in projects:
        nums = [s for rm in p["rooms"] for s in rm["slides"]]
        cnt = Counter(pic["md5"] for s in nums for pic in by_slide[s]["pictures"] if any(pic["crop"]))
        if not cnt:
            continue
        md5, n = cnt.most_common(1)[0]
        if n < 2:          # the plan is the cropped picture that recurs; 2 slides is enough in a short deck extract
            continue
        pic0 = next(pic for s in nums for pic in by_slide[s]["pictures"] if pic["md5"] == md5)
        p["plan"] = {"md5": md5, "file": pic0["file"], "px": pic0["px"], "slides": n}
        for s in nums:
            for pic in by_slide[s]["pictures"]:
                if pic["md5"] == md5:
                    pic["is_plan"] = True
                    pic["plan_rect_px"] = visible_rect(pic)
                    by_slide[s]["plan_overlays"] = overlays(by_slide[s], pic)


def visible_rect(pic):
    W, H = pic["px"]
    cl, ct, cr, cb = pic["crop"]
    return [round(cl * W, 1), round(ct * H, 1), round((1 - cr) * W, 1), round((1 - cb) * H, 1)]


def to_plan(pic, x, y):
    bx, by, bw, bh = pic["box"]
    x0, y0, x1, y1 = pic["plan_rect_px"]
    return [round(x0 + (x - bx) / bw * (x1 - x0), 1), round(y0 + (y - by) / bh * (y1 - y0), 1)]


def overlays(rec, pic):
    """Shapes and texts drawn on top of the plan picture, mapped into plan pixels."""
    bx, by, bw, bh = pic["box"]
    out = []
    for kind, items in (("shape", rec["shapes"]), ("text", rec["texts"])):
        for it in items:
            x, y, w, h = it["box"]
            cx, cy = x + w / 2, y + h / 2
            if bx <= cx <= bx + bw and by <= cy <= by + bh:
                a, b = to_plan(pic, x, y), to_plan(pic, x + w, y + h)
                o = {"type": kind, "rect_px": a + b}
                o.update({"kind": it["kind"]} if kind == "shape" else {"text": it["text"]})
                if it.get("points"):
                    o["points_px"] = [to_plan(pic, x, y) for x, y in it["points"]]
                out.append(o)
    return out


def read_pdf(deck, img_dir, seen):
    """A PDF deck: one 'slide' per page — text blocks with their boxes, embedded pictures, and a page render
    (pages/page-NN.png) for the agent and the app to look at. No drawn overlays (a PDF has no shape objects)."""
    import pymupdf as fitz
    doc = fitz.open(deck)
    (img_dir.parent / "pages").mkdir(parents=True, exist_ok=True)
    slides, size = [], None
    for i, page in enumerate(doc, 1):
        W, H = page.rect.width, page.rect.height
        size = size or [W, H]
        rec = {"slide": i, "texts": [], "pictures": [], "shapes": [], "tables": [], "page_png": f"pages/page-{i:02d}.png"}
        page.get_pixmap(dpi=110).save(str(img_dir.parent / rec["page_png"]))
        for x0, y0, x1, y1, txt, *_ in page.get_text("blocks"):
            if txt.strip():
                rec["texts"].append({"text": txt.strip(), "box": [round(x0 / W, 4), round(y0 / H, 4), round((x1 - x0) / W, 4), round((y1 - y0) / H, 4)]})
        for info in page.get_images(full=True):
            try:
                pix = fitz.Pixmap(doc, info[0])
                if pix.n - pix.alpha >= 4: pix = fitz.Pixmap(fitz.csRGB, pix)
                blob = pix.tobytes("png")
            except Exception:
                continue
            md5 = hashlib.md5(blob).hexdigest()
            if md5 not in seen:
                (img_dir / f"{md5[:12]}.png").write_bytes(blob)
                seen[md5] = {"file": f"images/{md5[:12]}.png", "px": [pix.width, pix.height]}
            rects = page.get_image_rects(info[0]) or [page.rect]
            r = rects[0]
            rec["pictures"].append({"md5": md5, "file": seen[md5]["file"], "px": seen[md5]["px"],
                                    "box": [round(r.x0 / W, 4), round(r.y0 / H, 4), round(r.width / W, 4), round(r.height / H, 4)], "crop": [0, 0, 0, 0]})
        all_text = " \n ".join(t["text"] for t in rec["texts"])
        rec["urls"] = sorted({u.rstrip(").,") for u in URL.findall(all_text)})
        rec["dims_mm"] = [("DIA " if d[0] else "") + d[1] + ("000" if d[2].lower() == "m" and not d[0] else "") for d in DIM.findall(all_text)]
        rec["levels_mm"] = [int(m[0]) for m in LEVEL.findall(all_text)]
        slides.append(rec)
    return slides, size


def fallback_project(slides, name):
    """No cover slide matched (a deck from another office, a PDF): one project named after the file, one room
    holding every slide. The user then names the room in the Questions step."""
    for r in slides:
        if r.get("kind") in (None, "front"): r["kind"] = "spec" if r["texts"] else "view"
        r.setdefault("category", category(r))
    return [{"title": name, "cover_slide": slides[0]["slide"] if slides else 1, "room_list": [], "fallback": True,
             "rooms": [{"name": name, "id": norm(name) or "room", "slides": [r["slide"] for r in slides]}]}]


def main(deck, out):
    out = Path(out)
    (out / "images").mkdir(parents=True, exist_ok=True)
    seen = {}
    if str(deck).lower().endswith(".pdf"):
        slides, (W, H) = read_pdf(deck, out / "images", seen)
    else:
        prs = Presentation(deck)
        W, H = prs.slide_width, prs.slide_height
        slides = [read_slide(i, s, W, H, out / "images", seen) for i, s in enumerate(prs.slides, 1)]
    projects = segment(slides) or fallback_project(slides, Path(deck).stem)
    by_slide = {r["slide"]: r for r in slides}
    plan_images(projects, by_slide)
    deck_json = {"schema": "deck/1", "source": str(Path(deck).name), "slide_size_emu": [W, H],
                 "slide_count": len(slides), "projects": projects, "slides": slides}
    (out / "deck.json").write_text(json.dumps(deck_json, indent=1, ensure_ascii=False), encoding="utf-8")
    for p in projects:
        print(f"[{p['cover_slide']}] {p['title']}  plan={'yes ' + str(p['plan']['px']) if 'plan' in p else 'no'}")
        for rm in p["rooms"]:
            kinds = Counter(by_slide[s]["kind"] for s in rm["slides"])
            print(f"    {rm['name']:<24} slides {rm['slides'][0] if rm['slides'] else '-'}..{rm['slides'][-1] if rm['slides'] else '-'}  {dict(kinds)}")
    print(f"{len(slides)} slides, {len(seen)} unique images -> {out}")


if __name__ == "__main__":
    main(*sys.argv[1:3])
