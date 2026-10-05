"""Generate a project's item folders, family schemas and TASKS.csv from its modelling workbook. Pure Python (needs openpyxl).

    python tools/workbook_to_project.py <PROJECT_CODE> [path/to/workbook.xlsx]

Workbook sheets read (see projects/_TEMPLATE/PROJECT.md for the column contract):
  Items       Item_Code, Space, Category, Item (…), Family, Host, Qty, L_mm, W_mm, H_mm, Pos_Z_mm, Faces, Params_JSON, Fit_Note
  Placements  Item_Code, Instance, X_mm, Y_mm, Z_mm, Rot_Z_deg
  Families    Family, Default_H_mm, Default_Params_JSON, Construction recipe
  Rooms       Space_Code, Space, X0_mm, Y0_mm, X1_mm, Y1_mm
Existing families are never overwritten: missing parameters are ADDED to family.json (version patch bump + changelog line);
missing families are created as status "spec" skeletons that an agent must implement.
"""
import json, os, re, sys, glob, datetime
from _boot import boot
ROOT = boot()
from openpyxl import load_workbook
from core import tasks as T

MATERIAL_KEYS = {"top_material", "fabric", "finish", "material", "image", "carcass", "cushion", "cct", "lumens"}
FRONT = {"S": 0, "E": 90, "N": 180, "W": 270, "SE": 45, "NE": 135, "NW": 225, "SW": 315}
DECOR_LAST = {"ACCESSORY", "SOFT_SET", "JOINERY_INSERT"}

def slug(s, n=40):
    return re.sub(r"[^A-Za-z0-9]+", "_", s).strip("_")[:n]

def sheet_rows(wb, name):
    ws = wb[name]; hdr = [c.value for c in ws[1]]
    return [dict(zip(hdr, [c.value for c in r])) for r in ws.iter_rows(min_row=2) if r[0].value is not None]

def bump(v):
    a, b, c = (int(x) for x in v.split(".")); return f"{a}.{b}.{c + 1}"

def ensure_family(name, fam_row, sample_params, project):
    d = os.path.join(ROOT, "families", name); os.makedirs(d, exist_ok=True)
    fj = os.path.join(d, "family.json"); today = datetime.date.today().isoformat()
    if os.path.exists(fj):
        spec = json.load(open(fj, encoding="utf-8")); added = []
        for k, v in sample_params.items():
            if k not in spec["parameters"]:
                spec["parameters"][k] = {"value": v, "description": f"added by project {project}"}; added.append(k)
        if added:
            spec["version"] = bump(spec["version"]); json.dump(spec, open(fj, "w", encoding="utf-8"), indent=1, ensure_ascii=False)
            with open(os.path.join(d, "CHANGELOG.md"), "a", encoding="utf-8") as f:
                f.write(f"- {today} v{spec['version']} — parameters added by {project}: {', '.join(added)}\n")
        return spec, ("extended" if added else "ok")
    defaults = json.loads(fam_row.get("Default_Params_JSON") or "{}") if fam_row else {}
    params = {k: {"value": v} for k, v in defaults.items() if k not in MATERIAL_KEYS}
    for k in ("L", "W", "H"):
        params.setdefault(k, {"value": (fam_row or {}).get("Default_H_mm", 500) if k == "H" else 500, "min": 1, "max": 30000, "unit": "mm"})
    for k, v in sample_params.items():
        params.setdefault(k, {"value": v, "description": f"added by project {project}"})
    spec = {"family": name, "version": "0.1.0", "status": "spec", "units": "mm", "front": "-Y", "host": "floor",
            "pivot": "bottom-centre", "recipe": (fam_row or {}).get("Construction recipe", "TODO: describe the geometry"), "parameters": params, "test_values": {}}
    json.dump(spec, open(fj, "w", encoding="utf-8"), indent=1, ensure_ascii=False)
    open(os.path.join(d, "build.py"), "w", encoding="utf-8").write(SKELETON.format(name=name))
    open(os.path.join(d, "CHANGELOG.md"), "w", encoding="utf-8").write(f"- {today} v0.1.0 — created as spec by {project}\n")
    return spec, "created"

SKELETON = '''"""Family {name} — generator (STATUS: spec, not implemented yet).
Contract: build(p, name, coll) -> root object at the origin, Z up, front toward -Y, pivot per family.json.
p = resolved parameters in mm (L, W, H + family parameters). Geometry only: no materials.
Implement from the recipe in family.json using core.geom; keep parts as separate child objects.
"""
from core import geom

def build(p, name, coll):
    raise NotImplementedError("Family {name} is a spec — implement families/{name}/build.py from family.json recipe")
'''

PROMPT = """# {code} — {item}

**Project:** {project} · **Space:** {space} · **Family:** {family} · **Host:** {host}

## Goal
Generate the mesh for this item (geometry only, no materials) with the `{family}` family, place every instance and report.

## Inputs
- `item.json` (this folder): size, parameters, placements, facing, notes.
- `families/{family}/family.json` (recipe + parameter schema) and `families/{family}/build.py` (generator).
- Core library: `core/` (geom, finish, place, checks, report). Read `AGENTS.md` first.

## Steps
1. Check `families/REGISTRY.json`: if `{family}` is `spec`, implement `families/{family}/build.py` from the recipe (core.geom only), add `test_values`, bump the family version, update its CHANGELOG and REGISTRY, run the family verify task.
2. In Blender run: `exec(compile(open(P).read(), P, "exec"), {{"__file__": P, "__name__": "__main__"}})` with `P` = full path of this folder's `build.py`.
3. Read `report.md`: status must be `ok`; fix non-manifold parts, triangle budget or overlaps in the FAMILY (not by hand), then rerun.
4. Update `TASKS.csv` (task `P2-ITEM-{code}`) and append a line to the project `REPORT.md` (done automatically by the task runner).

## Acceptance
Closed manifold solids (open edges allowed only for cards / curtains / sheets: `open_mesh_ok`), modifiers baked, scale 1, UVs present, triangles ≤ {budget}, instances = {n_inst}, names `{prefix}{code_u}_NN`, no unexplained overlaps.

## Size & placement (from item.json)
L × W × H = {L} × {W} × {H} mm · front faces {face} (rot_z {rot}°)
{places}
{note}
"""

BUILD = '''"""Build item {code} — {item}
Run inside Blender:
    P = r"<full path to this file>"; exec(compile(open(P).read(), P, "exec"), {{"__file__": P, "__name__": "__main__"}})
"""
import os, sys
HERE = os.path.dirname(os.path.abspath(__file__))
d = HERE
while not os.path.exists(os.path.join(d, "PIPELINE_ROOT")):
    d = os.path.dirname(d)
if d not in sys.path:
    sys.path.insert(0, d)
from tools.run_item import run_item
result = run_item(os.path.join(HERE, "item.json"))
print(result["status"], result.get("tris"), "tris", result["instances"], "instances", result["errors"] or "")
'''

def main(project, wb_path=None):
    pdir = os.path.join(ROOT, "projects", project)
    proj = json.load(open(os.path.join(pdir, "project.json"), encoding="utf-8"))
    wb_path = wb_path or glob.glob(os.path.join(pdir, "workbook", "*.xlsx"))[0]
    wb = load_workbook(wb_path, data_only=True)
    items = sheet_rows(wb, "Items"); places = sheet_rows(wb, "Placements")
    fams = {r["Family"]: r for r in sheet_rows(wb, "Families")}
    rooms = {r["Space"]: r for r in sheet_rows(wb, "Rooms")}
    pl_by = {}
    for p in places:
        pl_by.setdefault(p["Item_Code"], []).append(p)
    fam_status, fam_params = {}, {}
    for it in items:
        prm = {k: v for k, v in json.loads(it["Params_JSON"] or "{}").items() if k not in MATERIAL_KEYS}
        fam_params.setdefault(it["Family"], {}).update({k: v for k, v in prm.items() if k not in fam_params.get(it["Family"], {})})
    for name, sample in fam_params.items():
        spec, how = ensure_family(name, fams.get(name), sample, project); fam_status[name] = (spec["status"], how)
    reg_path = os.path.join(ROOT, "families", "REGISTRY.json")
    reg = json.load(open(reg_path)) if os.path.exists(reg_path) else {}
    for name in fam_params:
        spec = json.load(open(os.path.join(ROOT, "families", name, "family.json"), encoding="utf-8"))
        reg[name] = {"version": spec["version"], "status": spec["status"], "host": spec.get("host"), "projects": sorted(set(reg.get(name, {}).get("projects", []) + [project]))}
    json.dump(dict(sorted(reg.items())), open(reg_path, "w"), indent=1)
    rows = [dict(task_id="P0-SETUP", phase="0 setup", type="selftest", target="core", status="todo", notes="core selftest in Blender; open the project base .blend first")]
    for name in sorted(fam_params):
        st = reg[name]["status"]
        rows.append(dict(task_id=f"P1-FAM-{name}", phase="1 families", type="family_verify" if st in ("implemented", "tested") else "family_implement",
                         target=name, family=name, depends_on="P0-SETUP", status="todo", notes=f"family {st}"))
    space_items = {}
    furniture_tasks = {}
    order = sorted(items, key=lambda r: (r["Space"], r["Family"] in DECOR_LAST, r["Item_Code"]))
    for it in order:
        code = it["Item_Code"]; folder = f"{code}_{slug(it['Item (rev 3 wording)'] if 'Item (rev 3 wording)' in it else it['Item'])}"
        idir = os.path.join(pdir, "items", folder); os.makedirs(idir, exist_ok=True)
        name_txt = it.get("Item (rev 3 wording)") or it.get("Item")
        face = it.get("Faces") or "—"; rot = FRONT.get(face, None)
        pls = []
        for p in pl_by.get(code, []):
            z = p["Z_mm"]; z = z if isinstance(z, (int, float)) else "host"
            pls.append(dict(inst=int(p["Instance"]), x=p["X_mm"], y=p["Y_mm"], z=z, rot=rot if rot is not None else (p["Rot_Z_deg"] or 0)))
        room = rooms.get(it["Space"]); host = it["Host"]
        prm = {k: v for k, v in json.loads(it["Params_JSON"] or "{}").items() if k not in MATERIAL_KEYS}
        open_ok = it["Family"] in ("PLANTER_TREE", "PLANT_SMALL", "CURTAIN", "ACCESSORY", "SOFT_SET", "RUG", "WALL_PANEL_OGEE", "FINISH_WALL", "FINISH_FLOOR", "LIGHT_LED_STRIP", "MIRROR")
        budget = {"F": 60000}.get(it["Category"][0], 15000) if it["Family"] not in ("PLANTER_TREE",) else 120000
        ij = dict(code=code, project=project, space=it["Space"], space_code=code.split("-")[0], category=it["Category"], item=name_txt, family=it["Family"],
                  host=host, qty=it["Qty"], size=dict(L=it["L_mm"], W=it["W_mm"], H=it["H_mm"]), params=prm, face=face, rot_z=rot if rot is not None else 0,
                  placements=pls, room=[room["X0_mm"], room["Y0_mm"], room["X1_mm"], room["Y1_mm"]] if room else None, note=it.get("Fit_Note") or "",
                  tri_budget=budget, open_mesh_ok=open_ok, status="todo")
        json.dump(ij, open(os.path.join(idir, "item.json"), "w", encoding="utf-8"), indent=1, ensure_ascii=False)
        open(os.path.join(idir, "build.py"), "w", encoding="utf-8").write(BUILD.format(code=code, item=name_txt))
        places_txt = "\n".join(f"- #{p['inst']:02d} X {p['x']} · Y {p['y']} · Z {('FFL+' + str(p['z'])) if p['z'] != 'host' else 'host surface'} · rot {p['rot']}°" for p in pls) or "- room-wide (built in plan coordinates)"
        open(os.path.join(idir, "prompt.md"), "w", encoding="utf-8").write(PROMPT.format(code=code, item=name_txt, project=project, space=it["Space"], family=it["Family"], host=host,
            budget=budget, n_inst=len(pls) or 1, code_u=code.replace("-", "_"), prefix=proj.get("name_prefix", "SM_"), L=it["L_mm"], W=it["W_mm"], H=it["H_mm"], face=face, rot=ij["rot_z"], places=places_txt,
            note=("\n**Fit note:** " + ij["note"]) if ij["note"] else ""))
        if not os.path.exists(os.path.join(idir, "report.md")):
            open(os.path.join(idir, "report.md"), "w", encoding="utf-8").write(f"# {code} — build report\n\n- **Status:** not built yet\n")
        deps = [f"P1-FAM-{it['Family']}"]
        if it["Family"] in DECOR_LAST or host == "host":
            deps += furniture_tasks.get(it["Space"], [])
        tid = f"P2-ITEM-{code}"
        skip = "no geometry" in json.dumps(prm).lower()
        if skip:
            ij["status"] = "skipped"; json.dump(ij, open(os.path.join(idir, "item.json"), "w", encoding="utf-8"), indent=1, ensure_ascii=False)
        rows.append(dict(task_id=tid, phase="2 items", type="item_build", target=code, family=it["Family"], space=it["Space"], depends_on=";".join(deps),
                         status="skipped" if skip else "todo", notes="no geometry — skipped" if skip else "",
                         report=os.path.relpath(os.path.join(idir, "report.md"), pdir)))
        space_items.setdefault(it["Space"], []).append(tid)
        if it["Category"] == "Furniture" and host == "floor":
            furniture_tasks.setdefault(it["Space"], []).append(tid)
    for sp, tids in space_items.items():
        sc = tids[0].split("-")[2]
        rows.append(dict(task_id=f"P3-QC-{sc}", phase="3 qc", type="qc_space", target=f"{proj['collection_prefix']}{sc}", space=sp, depends_on=";".join(tids), status="todo"))
    qc = [r["task_id"] for r in rows if r["phase"] == "3 qc"]
    rows.append(dict(task_id="P4-EXPORT", phase="4 delivery", type="export", target="fbx", depends_on=";".join(qc), status="todo"))
    rows.append(dict(task_id="P4-REPORT", phase="4 delivery", type="report", target="REPORT.md", depends_on="P4-EXPORT", status="todo"))
    T.save(pdir, rows)
    print(f"{project}: {len(items)} items, {len(fam_params)} families ({sum(1 for s in fam_status.values() if s[1] == 'created')} created, "
          f"{sum(1 for s in fam_status.values() if s[1] == 'extended')} extended), {len(rows)} tasks")

if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else None)
